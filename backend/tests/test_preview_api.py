"""Phase 14 — text-preview endpoint tests.

Covers:
  * Plain-text extraction round-trip
  * Truncation at 100 KB
  * DOCX extraction (generated in-process via python-docx)
  * PDF extraction (hand-crafted minimal PDF — reportlab is not a dep)
  * Owner short-circuit
  * RBAC: non-owner non-admin denied while classification is pending
  * 415 for binary mime types
  * Audit row written with `file.previewed`
"""

from __future__ import annotations

import io
import uuid

import docx
from sqlalchemy import select

from app.models.audit import AuditLog
from app.models.classification import (
    Classification,
    ClassificationStatus,
    ClassificationTier,
)
from app.models.file import File
from app.services import preview

# Hand-crafted minimal PDF that pdfplumber can read. Built once so the
# test only pays the parse cost once per test.
_MINIMAL_PDF = b"""%PDF-1.4
1 0 obj
<< /Type /Catalog /Pages 2 0 R >>
endobj
2 0 obj
<< /Type /Pages /Kids [3 0 R] /Count 1 >>
endobj
3 0 obj
<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>
endobj
4 0 obj
<< /Length 55 >>
stream
BT /F1 24 Tf 100 700 Td (Hello Preview World) Tj ET
endstream
endobj
5 0 obj
<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>
endobj
xref
0 6
0000000000 65535 f
0000000009 00000 n
0000000054 00000 n
0000000097 00000 n
0000000196 00000 n
0000000302 00000 n
trailer
<< /Size 6 /Root 1 0 R >>
startxref
369
%%EOF
"""


def _build_docx_bytes() -> bytes:
    buf = io.BytesIO()
    document = docx.Document()
    document.add_paragraph("Hello DOCX preview")
    document.add_paragraph("Second paragraph with content")
    document.save(buf)
    return buf.getvalue()


def _upload(client, csrf_token, filename, content, content_type):
    return client.post(
        "/api/files",
        data={"file": (io.BytesIO(content), filename, content_type)},
        headers={"X-CSRF-Token": csrf_token},
        content_type="multipart/form-data",
    )


# --- Plain text ------------------------------------------------------------


def test_preview_text_file_returns_full_content(
    client, db_session, logged_in_admin, csrf_token
) -> None:
    body = b"hello\nworld\nthis is a text preview"
    up = _upload(client, csrf_token, "note.txt", body, "text/plain")
    file_id = up.json["id"]

    response = client.get(f"/api/files/{file_id}/preview")
    assert response.status_code == 200, response.data
    payload = response.json
    assert payload["extractor"] == "text"
    assert payload["truncated"] is False
    assert payload["text"] == body.decode("utf-8")
    assert payload["mime_type"] == "text/plain"


def test_preview_writes_audit_row(
    client, db_session, logged_in_admin, csrf_token
) -> None:
    up = _upload(client, csrf_token, "x.txt", b"hi", "text/plain")
    file_id = up.json["id"]

    client.get(f"/api/files/{file_id}/preview")

    audits = db_session.scalars(
        select(AuditLog).where(AuditLog.action_type == "file.previewed")
    ).all()
    assert len(audits) == 1
    meta = audits[0].event_metadata
    assert meta["extractor"] == "text"
    assert meta["filename"] == "x.txt"
    # Preview path must NOT double-emit a `file.downloaded` row.
    download_audits = db_session.scalars(
        select(AuditLog).where(AuditLog.action_type == "file.downloaded")
    ).all()
    assert len(download_audits) == 0


def test_preview_truncates_at_100kb(
    client, db_session, logged_in_admin, csrf_token
) -> None:
    # 200 KB of text — must be truncated to ~100 KB
    body = b"x" * (200 * 1024)
    up = _upload(client, csrf_token, "big.txt", body, "text/plain")
    file_id = up.json["id"]

    response = client.get(f"/api/files/{file_id}/preview")
    assert response.status_code == 200
    payload = response.json
    assert payload["truncated"] is True
    assert payload["preview_bytes"] <= preview.PREVIEW_BYTE_LIMIT
    assert payload["original_bytes"] == 200 * 1024
    assert len(payload["text"]) <= preview.PREVIEW_BYTE_LIMIT


def test_preview_treats_csv_as_text_even_with_octet_stream_mime(
    client, db_session, logged_in_admin, csrf_token
) -> None:
    """Some browsers upload .csv with mime=application/octet-stream. The
    extension fallback should still render it as text."""
    body = b"a,b,c\n1,2,3\n"
    up = _upload(client, csrf_token, "data.csv", body, "application/octet-stream")
    file_id = up.json["id"]
    response = client.get(f"/api/files/{file_id}/preview")
    assert response.status_code == 200
    assert response.json["extractor"] == "text"
    assert "1,2,3" in response.json["text"]


# --- PDF -------------------------------------------------------------------


def test_preview_pdf(client, db_session, logged_in_admin, csrf_token) -> None:
    up = _upload(client, csrf_token, "tiny.pdf", _MINIMAL_PDF, "application/pdf")
    file_id = up.json["id"]

    response = client.get(f"/api/files/{file_id}/preview")
    assert response.status_code == 200, response.data
    payload = response.json
    assert payload["extractor"] == "pdf"
    assert "Hello Preview World" in payload["text"]


# --- DOCX ------------------------------------------------------------------


def test_preview_docx(client, db_session, logged_in_admin, csrf_token) -> None:
    docx_bytes = _build_docx_bytes()
    up = _upload(
        client,
        csrf_token,
        "report.docx",
        docx_bytes,
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
    file_id = up.json["id"]

    response = client.get(f"/api/files/{file_id}/preview")
    assert response.status_code == 200, response.data
    payload = response.json
    assert payload["extractor"] == "docx"
    assert "Hello DOCX preview" in payload["text"]
    assert "Second paragraph with content" in payload["text"]


# --- Unsupported mime ------------------------------------------------------


def test_preview_unsupported_binary_returns_415(
    client, db_session, logged_in_admin, csrf_token
) -> None:
    # An arbitrary binary payload with a .bin extension — no extractor.
    body = b"\x00\x01\x02\x03\xff\xfe\xfd"
    up = _upload(client, csrf_token, "blob.bin", body, "application/octet-stream")
    file_id = up.json["id"]

    response = client.get(f"/api/files/{file_id}/preview")
    assert response.status_code == 415
    assert response.json["error"] == "unsupported_preview"


def test_preview_corrupt_pdf_returns_422(
    client, db_session, logged_in_admin, csrf_token
) -> None:
    """Mime says PDF but body is garbage → extractor crashes → 422."""
    up = _upload(
        client,
        csrf_token,
        "broken.pdf",
        b"not really a pdf",
        "application/pdf",
    )
    file_id = up.json["id"]
    response = client.get(f"/api/files/{file_id}/preview")
    assert response.status_code == 422
    assert response.json["error"] == "preview_failed"


# --- Access control --------------------------------------------------------


def test_preview_requires_auth(client, csrf_token) -> None:
    response = client.get(f"/api/files/{uuid.uuid4()}/preview")
    assert response.status_code == 401


def test_preview_not_found(client, db_session, logged_in_admin) -> None:
    response = client.get(f"/api/files/{uuid.uuid4()}/preview")
    assert response.status_code == 404


def test_preview_non_owner_denied_while_classification_pending(
    client, db_session, logged_in_user, seeded_admin
) -> None:
    """Same RBAC rule as `download`: while classification is pending, only
    the owner can read the file. A regular user trying to preview another
    user's file gets 403 with `classification_pending`.
    """
    other_file = File(
        owner_id=seeded_admin.id,
        original_filename="theirs.txt",
        storage_key=uuid.uuid4(),
        mime_type="text/plain",
        size_bytes=10,
        sha256="0" * 64,
        wrapped_dek=b"\x00" * 60,
        nonce=b"\x00" * 12,
        auth_tag=b"\x00" * 16,
        kek_version=1,
    )
    db_session.add(other_file)
    db_session.commit()
    response = client.get(f"/api/files/{other_file.id}/preview")
    assert response.status_code == 403
    assert response.json["reason"] == "classification_pending"


def test_preview_non_owner_allowed_after_public_classification(
    client, db_session, logged_in_user, seeded_admin, csrf_token
) -> None:
    """Once the admin's file is classified PUBLIC, a regular user's policy
    (allowed_max_tier=confidential) lets the preview through.
    """
    # Admin uploads a file via the real upload flow so encryption etc. is real.
    # We can't use logged_in_admin here (would log out logged_in_user). Insert
    # the file by hand + a COMPLETED classification at PUBLIC tier.
    plaintext = b"public file body for preview"
    from app.services import crypto, storage

    blob = crypto.encrypt_file(plaintext)
    storage_key = uuid.uuid4()
    storage.put(
        str(storage_key),
        io.BytesIO(blob.ciphertext),
        length=len(blob.ciphertext),
        content_type="application/octet-stream",
    )
    file_row = File(
        owner_id=seeded_admin.id,
        original_filename="public.txt",
        storage_key=storage_key,
        mime_type="text/plain",
        size_bytes=len(plaintext),
        sha256="a" * 64,
        wrapped_dek=blob.wrapped_dek,
        nonce=blob.nonce,
        auth_tag=blob.auth_tag,
        kek_version=blob.kek_version,
    )
    db_session.add(file_row)
    db_session.flush()
    db_session.add(
        Classification(
            file_id=file_row.id,
            version=1,
            status=ClassificationStatus.COMPLETED,
            final_tier=ClassificationTier.PUBLIC,
            zero_shot_tier=ClassificationTier.PUBLIC,
            classifier_version="test",
        )
    )
    db_session.commit()

    response = client.get(f"/api/files/{file_row.id}/preview")
    assert response.status_code == 200
    assert response.json["text"] == plaintext.decode()
