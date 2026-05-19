"""Text extraction for the classifier.

Each extractor opens a file and returns its textual contents as a single
string. Lifted verbatim from classify_sensitivity.py (root) so behavior
matches the standalone CLI. Imports are kept lazy per format to avoid
pulling every dependency on every call.
"""

from __future__ import annotations

from collections.abc import Callable
from pathlib import Path


def extract_pdf(path: Path) -> str:
    import pdfplumber

    with pdfplumber.open(path) as pdf:
        return "\n".join((p.extract_text() or "") for p in pdf.pages)


def extract_docx(path: Path) -> str:
    import docx

    return "\n".join(p.text for p in docx.Document(path).paragraphs if p.text)


def extract_pptx(path: Path) -> str:
    from pptx import Presentation

    parts: list[str] = []
    for slide in Presentation(path).slides:
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text:
                parts.append(shape.text)
    return "\n".join(parts)


def extract_text(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="ignore")


def extract_xlsx(path: Path) -> str:
    """Extract text from Excel .xlsx workbooks (all sheets, all cells)."""
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    parts: list[str] = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None]
            if cells:
                parts.append(" | ".join(cells))
    wb.close()
    return "\n".join(parts)


def extract_csv(path: Path) -> str:
    import csv

    parts: list[str] = []
    text = path.read_text(encoding="utf-8", errors="ignore")
    dialect = csv.Sniffer().sniff(text[:4096]) if text[:4096].strip() else None
    for row in csv.reader(text.splitlines(), dialect=dialect or "excel"):
        cells = [c for c in row if c.strip()]
        if cells:
            parts.append(" | ".join(cells))
    return "\n".join(parts)


def extract_rtf(path: Path) -> str:
    from striprtf.striprtf import rtf_to_text

    return rtf_to_text(path.read_text(encoding="utf-8", errors="ignore"))


def extract_html(path: Path) -> str:
    from html.parser import HTMLParser

    class _TextExtractor(HTMLParser):
        def __init__(self) -> None:
            super().__init__()
            self._parts: list[str] = []
            self._skip = False
            self._skip_tags = {"script", "style", "head", "meta", "link"}

        def handle_starttag(self, tag: str, _attrs) -> None:  # type: ignore[no-untyped-def]
            if tag.lower() in self._skip_tags:
                self._skip = True

        def handle_endtag(self, tag: str) -> None:
            if tag.lower() in self._skip_tags:
                self._skip = False

        def handle_data(self, data: str) -> None:
            if not self._skip and data.strip():
                self._parts.append(data.strip())

    parser = _TextExtractor()
    parser.feed(path.read_text(encoding="utf-8", errors="ignore"))
    return "\n".join(parser._parts)


def extract_json(path: Path) -> str:
    import json

    def _walk(obj):  # type: ignore[no-untyped-def]
        if isinstance(obj, str):
            return [obj]
        if isinstance(obj, dict):
            parts: list[str] = []
            for k, v in obj.items():
                parts.append(str(k))
                parts.extend(_walk(v))
            return parts
        if isinstance(obj, list):
            parts = []
            for item in obj:
                parts.extend(_walk(item))
            return parts
        return [str(obj)] if obj is not None else []

    data = json.loads(path.read_text(encoding="utf-8", errors="ignore"))
    return "\n".join(_walk(data))


def extract_yaml(path: Path) -> str:
    import yaml

    def _walk(obj, prefix: str = ""):  # type: ignore[no-untyped-def]
        if isinstance(obj, dict):
            parts: list[str] = []
            for k, v in obj.items():
                parts.extend(_walk(v, f"{prefix}{k}: "))
            return parts
        if isinstance(obj, list):
            parts = []
            for item in obj:
                parts.extend(_walk(item, prefix))
            return parts
        return [f"{prefix}{obj}"] if obj is not None else []

    data = yaml.safe_load(path.read_text(encoding="utf-8", errors="ignore"))
    return "\n".join(_walk(data))


# Extension → extractor. Plain-text variants (logs, source) reuse extract_text.
EXTRACTORS: dict[str, Callable[[Path], str]] = {
    # Documents
    ".pdf": extract_pdf,
    ".docx": extract_docx,
    ".doc": extract_docx,
    ".pptx": extract_pptx,
    ".rtf": extract_rtf,
    # Spreadsheets / tabular
    ".xlsx": extract_xlsx,
    ".csv": extract_csv,
    ".tsv": extract_csv,
    # Markup
    ".html": extract_html,
    ".htm": extract_html,
    # Data interchange
    ".json": extract_json,
    ".yaml": extract_yaml,
    ".yml": extract_yaml,
    # Plain text variants
    ".txt": extract_text,
    ".md": extract_text,
    ".rst": extract_text,
    ".log": extract_text,
    ".ini": extract_text,
    ".cfg": extract_text,
    ".conf": extract_text,
    ".env": extract_text,
    ".properties": extract_text,
    ".toml": extract_text,
    ".sql": extract_text,
}


class UnsupportedExtensionError(ValueError):
    """Raised when no extractor is registered for the file's extension."""


def get_text(path: Path) -> str:
    """Dispatch to the correct extractor based on file extension."""
    ext = path.suffix.lower()
    extractor = EXTRACTORS.get(ext)
    if extractor is None:
        raise UnsupportedExtensionError(f"unsupported extension: {ext}")
    return extractor(path)
